from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config_loader import ExcelConfig


class ExcelExporter:
    HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

    def __init__(self, output_dir: Path, excel_config: ExcelConfig):
        self.output_dir = output_dir
        self.config = excel_config
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export(self, rows: list[dict[str, Any]]) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = self.config.sheet_name

        headers = [c.header for c in self.config.columns]
        keys = [c.key for c in self.config.columns]

        ws.append(headers)
        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in rows:
            values = [self._coerce(row.get(k, "")) for k in keys]
            ws.append(values)

        for col_idx, key in enumerate(keys, start=1):
            max_len = max(
                [len(str(headers[col_idx - 1]))]
                + [len(str(r.get(key, ""))) for r in rows],
                default=12,
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

        ws.freeze_panes = "A2"

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = self.config.filename_template.format(timestamp=timestamp)
        output_path = self.output_dir / filename
        wb.save(output_path)
        return output_path

    @staticmethod
    def _coerce(value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, (int, float, str)):
            return value
        return str(value)
